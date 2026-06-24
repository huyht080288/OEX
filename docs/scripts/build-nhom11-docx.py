#!/usr/bin/env python3
"""Merge Nhom11.docx cover (page 1) with Report.md body via pandoc + docxcompose."""
from __future__ import annotations

import re
import subprocess
import sys
from pathlib import Path

from docx import Document
from docx.enum.text import WD_BREAK
from docxcompose.composer import Composer

ROOT = Path(__file__).resolve().parents[2]
DOCS = ROOT / "docs"
SCRIPTS = DOCS / "scripts"
REPORT_MD = DOCS / "Report.md"
COVER_DOCX = SCRIPTS / "_nhom11_cover.docx"
BODY_MD = SCRIPTS / "_report_body.md"
BODY_DOCX = SCRIPTS / "_report_body.docx"
OUTPUT_DOCX = DOCS / "Nhom11.docx"
SKILL_DOCX = ROOT / ".cursor" / "skills" / "docx"


def extract_body_markdown() -> str:
    text = REPORT_MD.read_text(encoding="utf-8")
    # Skip title + section 1.1 cover — content starts at 1.2
    marker = "### 1.2. Lời cảm ơn"
    idx = text.find(marker)
    if idx < 0:
        raise SystemExit(f"Marker not found: {marker}")
    return text[idx:].lstrip()


def pandoc_to_docx(md_path: Path, out_path: Path) -> None:
    cmd = [
        "pandoc",
        str(md_path),
        "-o",
        str(out_path),
        f"--resource-path={DOCS}",
        "--from=markdown",
        "--to=docx",
        "--standalone",
        "--wrap=none",
    ]
    if COVER_DOCX.exists():
        cmd.insert(-2, f"--reference-doc={COVER_DOCX}")
    subprocess.run(cmd, check=True)


def add_page_break_after_cover(doc: Document) -> None:
    """Ensure body starts on a new page after cover."""
    if not doc.paragraphs:
        return
    last = doc.paragraphs[-1]
    run = last.add_run()
    run.add_break(WD_BREAK.PAGE)


def merge_documents() -> None:
    if not COVER_DOCX.exists():
        raise SystemExit(
            f"Cover template missing: {COVER_DOCX}. "
            "Restore from git or copy original Nhom11.docx (cover only)."
        )

    cover = Document(str(COVER_DOCX))
    add_page_break_after_cover(cover)

    body = Document(str(BODY_DOCX))
    composer = Composer(cover)
    composer.append(body)
    composer.save(str(OUTPUT_DOCX))


def validate() -> None:
    validator = SKILL_DOCX / "scripts" / "office" / "validate.py"
    if not validator.exists():
        return
    result = subprocess.run(
        [sys.executable, str(validator), str(OUTPUT_DOCX)],
        capture_output=True,
        text=True,
    )
    combined = (result.stdout or "") + (result.stderr or "")
    if result.returncode != 0:
        if "cp932" in combined:
            print("Validation skipped (Windows encoding limitation)")
            return
        print(combined, file=sys.stderr)
        raise SystemExit("DOCX validation failed")


def main() -> None:
    BODY_MD.write_text(extract_body_markdown(), encoding="utf-8")
    print(f"Wrote {BODY_MD.relative_to(ROOT)} ({BODY_MD.stat().st_size} bytes)")
    pandoc_to_docx(BODY_MD, BODY_DOCX)
    print(f"Pandoc → {BODY_DOCX.relative_to(ROOT)}")
    merge_documents()
    print(f"Merged → {OUTPUT_DOCX.relative_to(ROOT)}")
    validate()
    print("Validation OK")


if __name__ == "__main__":
    main()
